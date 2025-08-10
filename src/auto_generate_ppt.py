from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData

from utils import remove_all_slides


def generate_ppt():
    # 创建 PPT
    prs = Presentation("../templates/MasterTemplate.pptx")
     # 清除模板中的所有幻灯片
    remove_all_slides(prs) 
    
    # 标题页
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "文本示例"
    # 图表页
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "图表示例"
    slide.shapes.add_picture("../images/forecast.png", Inches(1), Inches(2), width=Inches(6))
    # 表格页
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "表格示例"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "季度"
    table.cell(0, 1).text = "销量"
    table.cell(1, 0).text = "Q1"
    table.cell(1, 1).text = "300"
    # 保存
    prs.save("../pptoutputs/自动生成PPT示例.pptx")

if __name__ == "__main__":
    generate_ppt()